# -*- coding: utf-8 -*-
"""
§2.6 본문을 2개의 간결한 단락으로 교체 (현재 [194]-[195]가 잘못된 텍스트)
"""
import os, copy
from docx import Document
from lxml import etree

WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

folder = r'C:\Users\a00548169\OneDrive - ONEVIRTUALOFFICE\Desktop\자기주도적\검토용 논문'
path = os.path.join(folder, '(국문) 석사학위논문_임우현.docx')
doc = Document(path)

# §2.6 제목 찾기
title_idx = None
for i, p in enumerate(doc.paragraphs):
    if '2.6 전통 시계열 모형과의 비교' in p.text:
        title_idx = i
        break
assert title_idx is not None, '§2.6 제목을 찾을 수 없음'
print(f'Title at [{title_idx}]: {doc.paragraphs[title_idx].text}')

# 제목 다음 2개 단락 확인 (이전 스크립트에서 이미 7→2로 교체 완료)
# [194], [195]가 잘못된 유니코드 포함 → 삭제 후 재삽입
to_delete = [doc.paragraphs[title_idx + 1]._element,
             doc.paragraphs[title_idx + 2]._element]
print(f'Replacing: {doc.paragraphs[title_idx+1].text[:60]}')
print(f'Replacing: {doc.paragraphs[title_idx+2].text[:60]}')

for el in to_delete:
    el.getparent().remove(el)

NEW_BODY = [
    (
        "전통 시계열 모형(ARIMA, GARCH)은 가격·변동성의 자기이력에만 의존하는 "
        "선형 모형으로, 비선형 패턴 포착과 외부 정보(뉴스, 거시경제 변수) 편입이 "
        "구조적으로 불가능하다. "
        "본 연구는 이를 세 가지 관점에서 극복한다. "
        "첫째, Random Forest·XGBoost·LightGBM 등 트리 기반 앙상블과 BiLSTM으로 "
        "비선형 패턴을 학습한다. 실험 결과 GARCH(1,1)의 실현 변동성 R²는 "
        "미국 −0.213, 영국 −0.329인 반면, 트리 기반 모형은 로그 변동성 "
        "R² 0.830(미국)–0.932(영국)를 달성하였다. "
        "둘째, FinBERT 기반 768차원 임베딩·다중 헤드라인 16개 감성 통계량·금·유가·VIX 등 "
        "40여 개 금융 변수를 동시에 편입하여 외부 정보를 활용한다. "
        "Ablation 분석에서 AR_Only 대비 Full 모형의 AUC 증분은 "
        "미국 0.025·영국 0.053으로 통계적으로 유의하였다(미국 p<0.001, 영국 p=0.004). "
        "셋째, 동일한 파이프라인을 미국(SPY)과 영국(ISF.L)에 동시 적용하여 "
        "교차 시장 일반화를 검증하였다."
    ),
    (
        "이러한 세 가지 요소는 Walk-Forward 확장 윈도우(5-fold)·"
        "Block Bootstrap(10,000회)·Diebold-Mariano 검정을 결합한 "
        "엄밀한 out-of-sample 검증을 통해 실증되어, "
        "전통 시계열 모형 대비 본 연구의 방법론적 차별성을 확인한다."
    ),
]

title_para = doc.paragraphs[title_idx]
for text in reversed(NEW_BODY):
    ref = copy.deepcopy(title_para._element)
    for r in ref.findall(f'.//{{{WNS}}}r'):
        ref.remove(r)
    title_para._element.addnext(ref)
    r_el = etree.SubElement(ref, f'{{{WNS}}}r')
    t_el = etree.SubElement(r_el, f'{{{WNS}}}t')
    t_el.text = text
    t_el.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')

doc.save(path)
print('Saved.')

# 검증
doc2 = Document(path)
for i, p in enumerate(doc2.paragraphs):
    if '2.6 전통 시계열 모형' in p.text:
        for j in range(i, i + 4):
            print(f'[{j}] {doc2.paragraphs[j].text[:90]}')
        break
