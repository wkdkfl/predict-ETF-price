"""
검토용 논문에 v30 변경사항 적용
- 검토용 논문 == v29 (바이너리 동일 확인됨)
- _thesis_v30_update.py의 삽입 로직을 그대로 재활용
"""
import copy, shutil, sys
from docx import Document
from lxml import etree

TARGET = r"C:\Users\a00548169\OneDrive - ONEVIRTUALOFFICE\Desktop\자기주도적\검토용 논문\(국문) 석사학위논문_임우현.docx"

NEW_REFS = [
    (
        "[45] Du, K., Xing, F., Mao, R., & Cambria, E. (2024). "
        "Financial sentiment analysis: Techniques and applications. "
        "ACM Computing Surveys, 56(7), 1-36. "
        "DOI: https://doi.org/10.1145/3649451"
    ),
    (
        "[46] Zhang, C., Sjarif, N. N. A., & Ibrahim, R. (2024). "
        "Deep learning models for price forecasting of financial time series: "
        "A review of recent advancements: 2020-2022. "
        "WIREs Data Mining and Knowledge Discovery, 14(1), e1519. "
        "DOI: https://doi.org/10.1002/widm.1519"
    ),
    (
        "[47] Zhang, C., Zhang, Y., Cucuringu, M., & Qian, Z. (2024). "
        "Volatility forecasting with machine learning and intraday commonality. "
        "Journal of Financial Econometrics, 22(2), 492-530. "
        "DOI: https://doi.org/10.1093/jjfinec/nbad005"
    ),
    (
        "[48] Chen, W., Hussain, W., Cauteruccio, F., & Zhang, X. (2024). "
        "Deep learning for financial time series prediction: A state-of-the-art "
        "review of standalone and hybrid models. "
        "Computer Modeling in Engineering & Sciences, 139(1), 187-224. "
        "DOI: https://doi.org/10.32604/cmes.2023.031388"
    ),
    (
        "[49] Nie, Y., Kong, Y., Dong, X., Mulvey, J. M., Poor, H. V., "
        "Wen, Q., & Zohren, S. (2024). A survey of large language models "
        "for financial applications: Progress, prospects and challenges. "
        "arXiv preprint arXiv:2406.11903."
    ),
]

CH1_DIFFERENTIATION = (
    "본 연구의 접근법은 전통적 시계열 모형과 뚜렷이 구별된다. "
    "ARIMA는 선형 자기회귀·이동평균 구조만을 포착하며, GARCH(1,1)은 변동성 "
    "군집성을 모형화하되 외부 정보를 편입할 수 없다는 한계가 있다 [3, 40]. "
    "실제로 본 연구의 실험에서 GARCH(1,1)의 실현 변동성 R²는 미국 시장 "
    "−0.213, 영국 시장 −0.329로 나타나, 비선형 패턴과 외부 뉴스 정보를 반영하는 "
    "트리 기반 앙상블(log-R² 0.830–0.932)에 크게 미치지 못하였다. "
    "최근 Zhang 등 [47]은 ML 모형이 ARIMA·HAR 등 전통 모형 대비 변동성 예측에서 "
    "통계적으로 유의한 개선을 보임을 S&P 500에서 실증하였으며, "
    "Chen 등 [48]은 금융 시계열 예측에서 딥러닝 독립 모형 및 하이브리드 모형의 "
    "현황을 포괄적으로 비교하여 비선형 모형이 전통 통계 모형을 체계적으로 "
    "능가함을 보고하였다. 본 연구는 이러한 ML 기반 변동성 예측의 흐름에 FinBERT "
    "뉴스 감성 피처, 다중 헤드라인 집계, 교차 시장 검증이라는 세 가지 차별적 "
    "요소를 추가한다."
)

CH2_LLM_COMPARISON = (
    "한편, 최근 Du 등 [45]은 금융 감성 분석의 기법과 응용을 포괄적으로 조망하며, "
    "사전 기반·머신러닝·딥러닝·하이브리드·사전학습 모델의 다섯 범주를 체계적으로 "
    "비교하였다. 특히 Nie 등 [49]은 GPT, LLaMA 등 초대규모 언어모델(LLM)의 "
    "금융 응용을 65쪽에 걸쳐 서베이하면서, 감성 분석 과업에서 FinBERT와 같은 "
    "도메인 특화 모델이 범용 LLM 대비 비용 효율성과 재현성 측면에서 여전히 "
    "유력한 선택지임을 시사하였다. "
    "이러한 결과는 본 연구가 FinBERT를 선택한 근거를 한층 강화한다."
)

CH2_VOL_ML = (
    "최근 머신러닝 기반 변동성 예측 연구는 빠르게 확장되고 있다. "
    "Zhang 등 [47]은 ARIMA, HAR, OLS, LASSO, XGBoost, MLP, LSTM을 포함한 다수의 "
    "ML 모형을 S&P 500 구성 종목의 실현 변동성 예측에 적용하여, 신경망이 전통 모형 "
    "대비 통계적으로 유의한 개선을 달성함을 보고하였다. Zhang 등 [46]의 서베이는 "
    "2020–2022년 사이 발표된 DL 기반 가격 예측 모델을 체계적으로 비교하면서, "
    "Transformer·어텐션 메커니즘 기반 모형이 기존 LSTM·GRU 대비 예측 정확도를 "
    "유의하게 향상시키는 추세를 확인하였다. Chen 등 [48]은 독립형(standalone) 딥러닝과 "
    "하이브리드 모형을 38쪽에 걸쳐 비교하여, CNN-LSTM·Attention-LSTM 등 "
    "하이브리드 구조가 단일 모형 대비 일관된 성능 우위를 보임을 보고하였다. "
    "이들 연구는 본 연구가 채택한 트리 기반 앙상블 및 딥러닝 모형의 금융 시계열 "
    "예측 가치를 뒷받침한다."
)

CH2_HESTON = (
    "이와 관련하여, Heston과 Sinha [32]는 뉴스 기사의 원문(content) 자체가 "
    "단순 감성 점수보다 주가 수익률을 더 잘 예측함을 실증하였으며, 뉴스의 "
    "예측력이 출판 후 1–2일에 걸쳐 점진적으로 반영됨을 보고하였다. 이 결과는 "
    "감성 점수 외에 뉴스의 맥락 정보(임베딩)를 활용하는 본 연구의 접근법을 지지한다."
)

CH2_SEC26_TITLE = "2.6 전통 시계열 모형과의 비교 및 본 연구의 차별성"

CH2_SEC26_BODY_1 = (
    "본 절에서는 ARIMA·GARCH 등 전통적 시계열 모형의 한계를 체계적으로 정리하고, "
    "본 연구의 접근법이 이를 어떻게 극복하는지를 비교한다."
)

CH2_SEC26_BODY_2 = (
    "첫째, 비선형성의 한계이다. ARIMA는 선형 자기회귀 구조만을 포착하고, "
    "GARCH는 조건부 분산의 선형 점화식에 기반한다. 금융 시계열은 레버리지 효과, "
    "변동성 군집, 레짐 전환 등 비선형 동학을 내포하므로, 선형 모형만으로는 "
    "이를 충분히 반영하기 어렵다 [23, 40]. 본 연구는 Random Forest, XGBoost, "
    "LightGBM 등 트리 기반 앙상블과 BiLSTM, Transformer 등 딥러닝 모형을 "
    "활용하여 비선형 패턴을 학습한다. 실험 결과, 트리 기반 모형은 로그 변동성 "
    "R² 0.830(미국)–0.932(영국)를 달성한 반면, GARCH(1,1)의 실현 변동성 R²는 "
    "미국 −0.213, 영국 −0.329에 그쳤다. Chen 등 [48]의 포괄적 리뷰 역시 "
    "딥러닝 모형이 전통 통계 모형 대비 금융 시계열 예측에서 체계적으로 우월함을 "
    "확인하고 있다."
)

CH2_SEC26_BODY_3 = (
    "둘째, 외부 정보 편입의 한계이다. ARIMA와 GARCH는 가격·변동성의 자체 "
    "이력만을 입력으로 사용하며, 뉴스 텍스트나 거시경제 변수를 직접 포함할 수 없다. "
    "본 연구는 FinBERT 기반 768차원 임베딩과 3차원 감성 확률, 40여 개 금융 변수 "
    "(금, 유가, VIX, 환율 등)를 동시에 편입한다. Ablation 분석 결과, 방향성 예측에서 "
    "AR_Only 대비 Full 모형의 AUC 증분이 미국 0.025(0.724→0.749), "
    "영국 0.053(0.779→0.832)으로 나타났으며, Diebold-Mariano 검정에서도 "
    "뉴스 피처의 기여가 통계적으로 유의하였다(미국 p<0.001, 영국 p=0.004). "
    "Du 등 [45]의 포괄적 서베이 역시 금융 감성 분석이 변동성 예측과 방향성 분류 "
    "모두에서 유의한 부가 정보를 제공함을 확인한 바 있다."
)

CH2_SEC26_BODY_4 = (
    "셋째, 다중 헤드라인 집계이다. 전통 모형에는 텍스트 처리 기능이 없으므로, "
    "하루에 복수의 뉴스가 발생하더라도 이를 활용할 수 없다. 본 연구는 일자별 "
    "전체 헤드라인(평균 15–50건)을 FinBERT로 처리한 뒤, 16개 감성 통계량 "
    "(평균, 표준편차, 최대·최소, 분위수 등)과 평균 임베딩 벡터로 집계하여 "
    "날짜 수준의 텍스트 피처를 구축한다. Nie 등 [49]의 LLM 금융 서베이에서도 "
    "도메인 특화 언어모델의 감성 분류와 다중 문서 집계가 금융 예측에 유의한 "
    "증분 정보를 제공할 수 있음을 시사하고 있다."
)

CH2_SEC26_BODY_5 = (
    "넷째, 교차 시장 일반화이다. ARIMA·GARCH는 시장별로 별도의 모수를 추정해야 "
    "하며, 동일한 파이프라인이 다른 시장에서 유효한지에 대한 검증 기제가 부재하다. "
    "본 연구는 동일한 전처리, 피처 공학, 모델 아키텍처, 평가 프로토콜을 미국(SPY)과 "
    "영국(ISF.L)에 동시에 적용하여, 트리 기반 앙상블의 우위·변동성의 자기회귀 지배· "
    "방향성에 대한 뉴스의 유의한 증분 기여라는 동일한 패턴이 두 시장에서 재현됨을 확인하였다."
)

CH2_SEC26_BODY_6 = (
    "다섯째, 앙상블 다양성과 검증 엄밀성이다. 전통 모형이 단일 모수 추정에 의존하는 "
    "것과 달리, 본 연구의 Stacking 앙상블은 RF·XGBoost·LightGBM 세 기본 학습기의 "
    "예측을 Ridge 메타 학습기가 재결합한다. Walk-Forward 확장 윈도우(5-fold), "
    "Block Bootstrap(10,000회), Diebold-Mariano 검정을 결합하여 out-of-sample "
    "예측력의 통계적 유의성을 다각도로 검증하였다. Zhang 등 [47]과 Zhang 등 [46]이 "
    "공통적으로 강조하듯, walk-forward 기반의 엄밀한 out-of-sample 평가는 ML 금융 "
    "예측 연구의 신뢰성을 좌우하는 핵심 방법론적 요건이다."
)

CH2_SEC26_BODY_7 = (
    "이상의 다섯 가지 차별적 요소를 종합하면, 본 연구는 전통 시계열 모형의 선형성· "
    "단일 입력·단일 시장·단일 모형·제한적 검증이라는 한계를 동시에 극복하는 "
    "통합적 프레임워크를 제시한다."
)


WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

def find_para_idx(doc, keyword, start=0):
    for i, p in enumerate(doc.paragraphs):
        if i >= start and keyword in p.text:
            return i
    return None


def insert_paragraph_after(doc, idx, text):
    """Insert a new paragraph right after doc.paragraphs[idx], preserving style."""
    ref_para = doc.paragraphs[idx]
    new_p = copy.deepcopy(ref_para._element)
    for run_el in new_p.findall(f'.//{{{WNS}}}r'):
        new_p.remove(run_el)
    ref_para._element.addnext(new_p)
    r = etree.SubElement(new_p, f'{{{WNS}}}r')
    t = etree.SubElement(r, f'{{{WNS}}}t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return new_p


def main():
    doc = Document(TARGET)
    original_count = len(doc.paragraphs)
    print(f"Original paragraph count: {original_count}")

    # 1. 참고문헌 [45]-[49] 추가
    last_ref_idx = find_para_idx(doc, "[44] Krauss, C.")
    if last_ref_idx is None:
        last_ref_idx = find_para_idx(doc, "[44]")
    assert last_ref_idx is not None, "FAIL: Cannot find [44] reference"
    print(f"[1/6] 참고문헌 [44] at paragraph {last_ref_idx}")

    cursor = last_ref_idx
    for ref_text in NEW_REFS:
        insert_paragraph_after(doc, cursor, ref_text)
        cursor += 1
    print(f"  -> Inserted {len(NEW_REFS)} references ([45]-[49])")

    # 2. Chapter 1: 차별성 요약 단락
    ch1_idx = find_para_idx(doc, "본 연구의 학술적 기여는 다음과 같다")
    assert ch1_idx is not None, "FAIL: Cannot find 학술적 기여 paragraph"
    print(f"[2/6] Ch1 학술적 기여 at paragraph {ch1_idx}")
    insert_paragraph_after(doc, ch1_idx, CH1_DIFFERENTIATION)
    print("  -> Inserted differentiation summary")

    # 3. §2.3 FinBERT: LLM 비교
    llm_idx = find_para_idx(doc, "Lopez-Lira와 Tang [22]")
    if llm_idx is None:
        llm_idx = find_para_idx(doc, "LLM을 금융 감성")
    assert llm_idx is not None, "FAIL: Cannot find LLM paragraph in §2.3"
    print(f"[3/6] §2.3 LLM paragraph at {llm_idx}")
    insert_paragraph_after(doc, llm_idx, CH2_LLM_COMPARISON)
    print("  -> Inserted LLM comparison (Du[45], Nie[49])")

    # 4. §2.2: Heston & Sinha [32]
    heston_idx = find_para_idx(doc, "뉴스 헤드라인이 소셜 미디어 대비 노이즈가 적고")
    if heston_idx is None:
        heston_idx = find_para_idx(doc, "Bollen 등 [21]")
    assert heston_idx is not None, "FAIL: Cannot find §2.2 social media paragraph"
    print(f"[4/6] §2.2 social media paragraph at {heston_idx}")
    insert_paragraph_after(doc, heston_idx, CH2_HESTON)
    print("  -> Inserted Heston & Sinha [32]")

    # 5. §2.5: ML 변동성 최신 논문
    vol_idx = find_para_idx(doc, "변동성 예측은 금융계량경제학의 핵심")
    assert vol_idx is not None, "FAIL: Cannot find §2.5 volatility paragraph"
    print(f"[5/6] §2.5 volatility paragraph at {vol_idx}")
    insert_paragraph_after(doc, vol_idx, CH2_VOL_ML)
    print("  -> Inserted ML volatility research (Zhang[47], Zhang[46], Chen[48])")

    # 6. §2.6 새 서브섹션 (제목 + 7 본문 단락)
    sec25_end_idx = find_para_idx(doc, "본 연구는 이상의 선행 연구들이 각각 탐구해 온")
    assert sec25_end_idx is not None, "FAIL: Cannot find §2.5 closing paragraph"
    print(f"[6/6] §2.5 closing paragraph at {sec25_end_idx}")

    # 역순으로 삽입하여 올바른 순서 유지
    bodies = [
        CH2_SEC26_BODY_7,
        CH2_SEC26_BODY_6,
        CH2_SEC26_BODY_5,
        CH2_SEC26_BODY_4,
        CH2_SEC26_BODY_3,
        CH2_SEC26_BODY_2,
        CH2_SEC26_BODY_1,
        CH2_SEC26_TITLE,
    ]
    insert_before_idx = sec25_end_idx - 1
    for body_text in bodies:
        insert_paragraph_after(doc, insert_before_idx, body_text)
    print("  -> Inserted §2.6 (title + 7 body paragraphs)")

    # 저장
    doc.save(TARGET)
    final_count = len(Document(TARGET).paragraphs)
    print(f"\nSaved: {TARGET}")
    print(f"Paragraphs: {original_count} -> {final_count} (+{final_count - original_count})")

    # 검증
    verify = Document(TARGET)
    checks = {
        "[49] reference": any("[49]" in p.text for p in verify.paragraphs),
        "§2.6 title": any("2.6 전통 시계열 모형" in p.text for p in verify.paragraphs),
        "ARIMA differentiation": any("ARIMA는 선형 자기회귀" in p.text for p in verify.paragraphs),
        "Du[45] citation": any("Du 등 [45]" in p.text for p in verify.paragraphs),
        "Heston[32] citation": any("Heston과 Sinha [32]" in p.text for p in verify.paragraphs),
    }
    print("\nVerification:")
    all_ok = True
    for name, found in checks.items():
        status = "OK" if found else "MISSING"
        if not found:
            all_ok = False
        print(f"  {name}: {status}")

    if all_ok:
        print("\nAll checks passed!")
    else:
        print("\nWARNING: Some checks failed!")
        sys.exit(1)


if __name__ == "__main__":
    main()
