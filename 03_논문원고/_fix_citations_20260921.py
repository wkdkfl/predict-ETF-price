# -*- coding: utf-8 -*-
"""인용 표기 정리 (2026-09-21).

논문은 IEEE 번호식([n])으로 통일돼 있으나, 괄호 인용 11곳에만
'저자, 연도 [번호]' 형태의 옛 APA 잔재가 남아 있다. 번호가 이미 있으므로
저자·연도는 순수한 중복이다. 이를 번호식으로 정리한다.

  (EMH; Fama, 1970 [11])                      -> (EMH) [11]
  (Cont, 2001 [5]; Corsi, 2009 [6])           -> [5], [6]
  (Christoffersen & Diebold, 2006 [37])       -> [37]
  ...

서술형 인용('Tetlock [17]은 …')은 건드리지 않는다 — 문헌 검토에서 주어 역할을 하며
번호만 남기면 문장이 성립하지 않는다.

실행: python _fix_citations_20260921.py [--dry-run]
"""
from __future__ import annotations

import sys
from pathlib import Path

import docx

DOC = Path(__file__).resolve().parent / "(국문) 석사학위논문_임우현.docx"

# (찾을 문자열, 바꿀 문자열) — 문단 전체 텍스트 기준. 번호는 오름차순으로 정렬해 표기한다.
REPLACEMENTS = [
    ("효율적 시장가설(EMH; Fama, 1970 [11])과",
     "효율적 시장가설(EMH) [11]과"),
    ("군집성을 보이며(Cont, 2001 [5]; Corsi, 2009 [6]),",
     "군집성을 보이며 [5], [6],"),
    ("알려져 있다(Christoffersen & Diebold, 2006 [37]).",
     "알려져 있다 [37]."),
    ("축적되어 있다(Engle & Ng, 1993 [33]; Andersen et al., 2003 [32]; Corsi, 2009 [6]).",
     "축적되어 있다 [6], [32], [33]."),
    ("보고되어 왔다(Leitch & Tanner, 1991 [36]; Christoffersen & Diebold, 2006 [37]).",
     "보고되어 왔다 [36], [37]."),
    ("비대칭 변동성 패턴(Black, 1976 [34]; Engle & Ng, 1993 [33])이",
     "비대칭 변동성 패턴 [33], [34]이"),
]


def iter_paragraphs(document):
    """본문 + 표 셀 안의 모든 문단."""
    for p in document.paragraphs:
        yield p
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def replace_in_paragraph(par, old, new):
    """런(run) 서식을 보존하며 문단 텍스트의 old 를 new 로 치환. 치환 횟수를 돌려준다."""
    runs = par.runs
    if not runs:
        return 0
    n = 0
    while True:
        full = "".join(r.text for r in runs)
        at = full.find(old)
        if at < 0:
            return n
        end = at + len(old)
        # 각 런의 [시작, 끝) 오프셋
        pos, spans = 0, []
        for r in runs:
            spans.append((pos, pos + len(r.text)))
            pos += len(r.text)
        first = next(i for i, (a, b) in enumerate(spans) if b > at)
        placed = False
        for i, (a, b) in enumerate(spans):
            if b <= at or a >= end:
                continue                      # 매치와 무관한 런
            head = runs[i].text[: max(0, at - a)]
            tail = runs[i].text[max(0, end - a):]
            runs[i].text = head + (new if not placed and i == first else "") + tail
            placed = placed or i == first
        n += 1


def main():
    dry = "--dry-run" in sys.argv
    document = docx.Document(str(DOC))
    total = 0
    for old, new in REPLACEMENTS:
        hits = 0
        for par in iter_paragraphs(document):
            hits += replace_in_paragraph(par, old, new)
        total += hits
        status = "OK " if hits else "못찾음"
        print("%s %d회 | %s\n        -> %s" % (status, hits, old[:68], new[:68]))
    print("\n총 치환 %d회" % total)
    if dry:
        print("(--dry-run: 저장하지 않음)")
        return
    document.save(str(DOC))
    print("저장 완료: %s" % DOC.name)


if __name__ == "__main__":
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    main()
